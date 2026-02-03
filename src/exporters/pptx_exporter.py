"""
Exportador a formato PowerPoint (PPTX).
Genera presentaciones profesionales con análisis de datos.
"""
import os
from typing import Any, Dict, List
from datetime import datetime
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE


# Paleta de colores profesional
COLORS = {
    "primary": RGBColor(0, 82, 147),      # Azul corporativo
    "secondary": RGBColor(0, 128, 128),   # Teal
    "accent": RGBColor(255, 140, 0),      # Naranja
    "success": RGBColor(40, 167, 69),     # Verde
    "warning": RGBColor(255, 193, 7),     # Amarillo
    "danger": RGBColor(220, 53, 69),      # Rojo
    "dark": RGBColor(52, 58, 64),         # Gris oscuro
    "light": RGBColor(248, 249, 250),     # Gris claro
    "white": RGBColor(255, 255, 255),
    "black": RGBColor(0, 0, 0)
}


def export_to_pptx(
    data: Any, 
    analysis: Dict[str, Any],
    output_dir: str = "/app/outputs"
) -> str:
    """
    Exporta los datos a una presentación PowerPoint profesional.
    
    Args:
        data: Datos a incluir en la presentación
        analysis: Metadatos y análisis
        output_dir: Directorio de salida
    
    Returns:
        Ruta al archivo generado
    """
    # Crear presentación 16:9
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    analysis_type = analysis.get("type", "query")
    
    # Slide 1: Título
    _add_title_slide(prs, analysis_type)
    
    # Slides según tipo de análisis
    if analysis_type == "table_analysis":
        _add_table_analysis_slides(prs, data, analysis)
    elif analysis_type == "database_report":
        _add_database_report_slides(prs, data, analysis)
    else:
        _add_query_results_slides(prs, data, analysis)
    
    # Slide de insights
    insights = analysis.get("insights", [])
    if insights:
        _add_insights_slide(prs, insights)
    
    # Slide final
    _add_closing_slide(prs)
    
    # Guardar archivo
    os.makedirs(output_dir, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"db_analysis_{timestamp}.pptx"
    filepath = os.path.join(output_dir, filename)
    prs.save(filepath)
    
    return filepath


def _add_title_slide(prs: Presentation, analysis_type: str):
    """Agrega slide de título con diseño profesional."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    
    # Fondo con gradiente simulado (barra superior)
    header_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0), Inches(0),
        Inches(13.333), Inches(2.5)
    )
    header_shape.fill.solid()
    header_shape.fill.fore_color.rgb = COLORS["primary"]
    header_shape.line.fill.background()
    
    # Título principal
    title_box = slide.shapes.add_textbox(
        Inches(0.5), Inches(0.8),
        Inches(12.333), Inches(1.2)
    )
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "📊 Análisis de Base de Datos"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = COLORS["white"]
    p.alignment = PP_ALIGN.CENTER
    
    # Subtítulo según tipo
    subtitles = {
        "table_analysis": "Análisis de Estructura y Datos",
        "database_report": "Reporte General de Base de Datos",
        "query": "Resultados de Consulta SQL"
    }
    subtitle_text = subtitles.get(analysis_type, "Análisis de Datos")
    
    subtitle_box = slide.shapes.add_textbox(
        Inches(0.5), Inches(1.9),
        Inches(12.333), Inches(0.5)
    )
    tf = subtitle_box.text_frame
    p = tf.paragraphs[0]
    p.text = subtitle_text
    p.font.size = Pt(24)
    p.font.color.rgb = COLORS["light"]
    p.alignment = PP_ALIGN.CENTER
    
    # Fecha y hora
    date_box = slide.shapes.add_textbox(
        Inches(0.5), Inches(6.5),
        Inches(12.333), Inches(0.5)
    )
    tf = date_box.text_frame
    p = tf.paragraphs[0]
    p.text = f"Generado: {datetime.now().strftime('%d de %B de %Y, %H:%M')}"
    p.font.size = Pt(14)
    p.font.color.rgb = COLORS["dark"]
    p.alignment = PP_ALIGN.CENTER


def _add_query_results_slides(prs: Presentation, data: Any, analysis: Dict[str, Any]):
    """Agrega slides con resultados de query."""
    
    # Slide de resumen de query
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_slide_header(slide, "📋 Resumen de la Consulta")
    
    query_info = analysis.get("query_info", {})
    results_info = analysis.get("results_info", {})
    
    # Info boxes
    y_pos = 1.8
    
    # Tipo de query
    _add_info_box(slide, Inches(0.5), Inches(y_pos), Inches(4), Inches(1.2),
                  "Tipo de Query", query_info.get("type", "N/A"), COLORS["primary"])
    
    # Complejidad
    complexity = query_info.get("complexity", "N/A")
    complexity_color = {
        "simple": COLORS["success"],
        "moderate": COLORS["warning"],
        "complex": COLORS["accent"],
        "very_complex": COLORS["danger"]
    }.get(complexity, COLORS["dark"])
    
    _add_info_box(slide, Inches(4.7), Inches(y_pos), Inches(4), Inches(1.2),
                  "Complejidad", complexity.upper(), complexity_color)
    
    # Registros
    _add_info_box(slide, Inches(8.9), Inches(y_pos), Inches(4), Inches(1.2),
                  "Registros", f"{results_info.get('row_count', 0):,}", COLORS["secondary"])
    
    # Tablas involucradas
    tables = query_info.get("tables", [])
    if tables:
        tables_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(3.5),
            Inches(12.333), Inches(0.8)
        )
        tf = tables_box.text_frame
        p = tf.paragraphs[0]
        p.text = f"Tablas: {', '.join(tables)}"
        p.font.size = Pt(16)
        p.font.color.rgb = COLORS["dark"]
    
    # Columnas
    columns = results_info.get("columns", [])
    if columns:
        cols_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(4.3),
            Inches(12.333), Inches(1.5)
        )
        tf = cols_box.text_frame
        p = tf.paragraphs[0]
        p.text = f"Columnas ({len(columns)}): {', '.join(columns[:8])}"
        if len(columns) > 8:
            p.text += f" ... y {len(columns) - 8} más"
        p.font.size = Pt(14)
        p.font.color.rgb = COLORS["dark"]
    
    # Slides de datos
    if isinstance(data, list) and data:
        _add_data_table_slides(prs, data)


def _add_table_analysis_slides(prs: Presentation, data: Dict[str, Any], analysis: Dict[str, Any]):
    """Agrega slides de análisis de tabla."""
    table_name = data.get("table_name", "Unknown")
    
    # Slide de estructura
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_slide_header(slide, f"🗃️ Estructura: {table_name}")
    
    structure = data.get("structure", {})
    columns = structure.get("columns", [])
    
    if columns:
        # Crear tabla de estructura
        rows = min(len(columns), 12) + 1  # +1 para header
        table = slide.shapes.add_table(
            rows, 4,
            Inches(0.5), Inches(1.5),
            Inches(12.333), Inches(5)
        ).table
        
        # Estilo de header
        headers = ["Columna", "Tipo", "Nullable", "PK"]
        for i, header in enumerate(headers):
            cell = table.cell(0, i)
            cell.text = header
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLORS["primary"]
            p = cell.text_frame.paragraphs[0]
            p.font.bold = True
            p.font.color.rgb = COLORS["white"]
            p.font.size = Pt(12)
        
        # Datos
        for row_idx, col in enumerate(columns[:12]):
            table.cell(row_idx + 1, 0).text = col.get("name", "")[:25]
            table.cell(row_idx + 1, 1).text = str(col.get("type", ""))[:20]
            table.cell(row_idx + 1, 2).text = "Sí" if col.get("nullable") else "No"
            table.cell(row_idx + 1, 3).text = "✓" if col.get("primary_key") else ""
            
            for j in range(4):
                p = table.cell(row_idx + 1, j).text_frame.paragraphs[0]
                p.font.size = Pt(11)
    
    # Slide de estadísticas
    stats = data.get("statistics", {})
    if stats:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _add_slide_header(slide, f"📊 Estadísticas: {table_name}")
        
        _add_info_box(slide, Inches(2), Inches(2.5), Inches(4.5), Inches(2),
                      "Total de Filas", f"{stats.get('total_rows', 0):,}", COLORS["primary"])
        
        if stats.get("table_size"):
            _add_info_box(slide, Inches(7), Inches(2.5), Inches(4.5), Inches(2),
                          "Tamaño", stats.get("table_size"), COLORS["secondary"])


def _add_database_report_slides(prs: Presentation, data: Dict[str, Any], analysis: Dict[str, Any]):
    """Agrega slides de reporte de base de datos."""
    overview = data.get("database_overview", {})
    
    # Slide de overview
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_slide_header(slide, "🗄️ Vista General de la Base de Datos")
    
    _add_info_box(slide, Inches(4), Inches(2.5), Inches(5.333), Inches(2.5),
                  "Total de Tablas", str(overview.get("total_tables", 0)), COLORS["primary"])
    
    # Slide con lista de tablas
    table_analyses = data.get("table_analyses", [])
    if table_analyses:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _add_slide_header(slide, "📋 Tablas en la Base de Datos")
        
        rows = min(len(table_analyses), 10) + 1
        table = slide.shapes.add_table(
            rows, 3,
            Inches(1), Inches(1.5),
            Inches(11.333), Inches(5)
        ).table
        
        # Headers
        for i, header in enumerate(["Tabla", "Columnas", "Filas"]):
            cell = table.cell(0, i)
            cell.text = header
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLORS["primary"]
            p = cell.text_frame.paragraphs[0]
            p.font.bold = True
            p.font.color.rgb = COLORS["white"]
        
        # Datos
        for row_idx, ta in enumerate(table_analyses[:10]):
            table.cell(row_idx + 1, 0).text = ta.get("table_name", "")[:30]
            table.cell(row_idx + 1, 1).text = str(ta.get("columns", "N/A"))
            rows_val = ta.get("rows", "N/A")
            table.cell(row_idx + 1, 2).text = f"{rows_val:,}" if isinstance(rows_val, int) else str(rows_val)


def _add_data_table_slides(prs: Presentation, data: List[Dict[str, Any]]):
    """Agrega slides con tablas de datos."""
    if not data:
        return
    
    columns = list(data[0].keys())[:6]  # Máximo 6 columnas
    rows_per_slide = 8
    
    for batch_idx in range(0, min(len(data), 24), rows_per_slide):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        batch_end = min(batch_idx + rows_per_slide, len(data))
        _add_slide_header(slide, f"📝 Datos (Registros {batch_idx + 1} - {batch_end})")
        
        batch = data[batch_idx:batch_end]
        
        table = slide.shapes.add_table(
            len(batch) + 1, len(columns),
            Inches(0.3), Inches(1.4),
            Inches(12.733), Inches(5.5)
        ).table
        
        # Headers
        for i, col in enumerate(columns):
            cell = table.cell(0, i)
            cell.text = str(col)[:15]
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLORS["primary"]
            p = cell.text_frame.paragraphs[0]
            p.font.bold = True
            p.font.color.rgb = COLORS["white"]
            p.font.size = Pt(10)
        
        # Datos
        for row_idx, row in enumerate(batch):
            for col_idx, col in enumerate(columns):
                cell = table.cell(row_idx + 1, col_idx)
                value = row.get(col, "")
                cell.text = str(value)[:25] if value else ""
                p = cell.text_frame.paragraphs[0]
                p.font.size = Pt(9)


def _add_insights_slide(prs: Presentation, insights: List[str]):
    """Agrega slide de insights."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_slide_header(slide, "💡 Insights y Observaciones")
    
    y_pos = 1.8
    for insight in insights[:8]:
        insight_box = slide.shapes.add_textbox(
            Inches(0.8), Inches(y_pos),
            Inches(11.733), Inches(0.6)
        )
        tf = insight_box.text_frame
        p = tf.paragraphs[0]
        p.text = f"• {insight}"
        p.font.size = Pt(16)
        p.font.color.rgb = COLORS["dark"]
        y_pos += 0.65


def _add_closing_slide(prs: Presentation):
    """Agrega slide de cierre."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # Fondo
    bg_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0), Inches(0),
        Inches(13.333), Inches(7.5)
    )
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = COLORS["primary"]
    bg_shape.line.fill.background()
    
    # Texto
    text_box = slide.shapes.add_textbox(
        Inches(0.5), Inches(3),
        Inches(12.333), Inches(1.5)
    )
    tf = text_box.text_frame
    p = tf.paragraphs[0]
    p.text = "¿Preguntas?"
    p.font.size = Pt(48)
    p.font.bold = True
    p.font.color.rgb = COLORS["white"]
    p.alignment = PP_ALIGN.CENTER
    
    # Subtexto
    sub_box = slide.shapes.add_textbox(
        Inches(0.5), Inches(4.5),
        Inches(12.333), Inches(0.5)
    )
    tf = sub_box.text_frame
    p = tf.paragraphs[0]
    p.text = "Generado con MCP Database Analyzer"
    p.font.size = Pt(16)
    p.font.color.rgb = COLORS["light"]
    p.alignment = PP_ALIGN.CENTER


def _add_slide_header(slide, title: str):
    """Agrega header estándar a un slide."""
    # Barra de header
    header_bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(0), Inches(0),
        Inches(13.333), Inches(1.2)
    )
    header_bar.fill.solid()
    header_bar.fill.fore_color.rgb = COLORS["primary"]
    header_bar.line.fill.background()
    
    # Título
    title_box = slide.shapes.add_textbox(
        Inches(0.5), Inches(0.35),
        Inches(12.333), Inches(0.6)
    )
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = COLORS["white"]


def _add_info_box(slide, left, top, width, height, label: str, value: str, color: RGBColor):
    """Agrega una caja de información estilizada."""
    # Fondo
    box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        left, top, width, height
    )
    box.fill.solid()
    box.fill.fore_color.rgb = color
    box.line.fill.background()
    
    # Label
    label_box = slide.shapes.add_textbox(
        left, top + Inches(0.2),
        width, Inches(0.4)
    )
    tf = label_box.text_frame
    p = tf.paragraphs[0]
    p.text = label
    p.font.size = Pt(12)
    p.font.color.rgb = COLORS["white"]
    p.alignment = PP_ALIGN.CENTER
    
    # Value
    value_box = slide.shapes.add_textbox(
        left, top + Inches(0.5),
        width, Inches(0.6)
    )
    tf = value_box.text_frame
    p = tf.paragraphs[0]
    p.text = value
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = COLORS["white"]
    p.alignment = PP_ALIGN.CENTER
